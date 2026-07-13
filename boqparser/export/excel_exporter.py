"""Excel Exporter — multi-sheet XLSX with section formatting and validation highlighting."""

from __future__ import annotations

from pathlib import Path
from typing import Any, List


def _safe_float(val) -> Any:
    """Return float or None."""
    if val is None or val == "":
        return None
    try:
        return float(val)
    except (TypeError, ValueError):
        return None


class ExcelExporter:
    """Export BOQ parse result to multi-sheet XLSX using openpyxl."""

    def export(self, parse_result: Any, output_path: str) -> str:
        """Write to XLSX file. Returns absolute path."""
        try:
            import openpyxl
            from openpyxl.styles import (
                Font, PatternFill, Alignment, Border, Side
            )
        except ImportError:
            raise ImportError("openpyxl is required for Excel export. pip install openpyxl")

        wb = openpyxl.Workbook()
        wb.remove(wb.active)  # remove default sheet

        # ── Styles ────────────────────────────────────────────────────────────
        title_font   = Font(name="Calibri", bold=True, size=14, color="FFFFFF")
        header_font  = Font(name="Calibri", bold=True, size=10, color="FFFFFF")
        section_font = Font(name="Calibri", bold=True, size=10, color="003366")
        normal_font  = Font(name="Calibri", size=10)
        error_font   = Font(name="Calibri", size=10, color="CC0000")

        title_fill   = PatternFill("solid", fgColor="003366")
        header_fill  = PatternFill("solid", fgColor="1F4E79")
        section_fill = PatternFill("solid", fgColor="D6E4F0")
        subtotal_fill = PatternFill("solid", fgColor="FFF2CC")
        error_fill   = PatternFill("solid", fgColor="FFE0E0")
        alt_fill     = PatternFill("solid", fgColor="F5F9FF")

        thin = Side(border_style="thin", color="BFBFBF")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        center_align = Alignment(horizontal="center", vertical="center")
        left_align   = Alignment(horizontal="left", vertical="center", wrap_text=True)
        right_align  = Alignment(horizontal="right", vertical="center")

        col_widths = [10, 45, 12, 8, 12, 14]  # item_no, desc, qty, unit, rate, amount
        headers = ["Item No.", "Description", "Quantity", "Unit", "Rate (₹)", "Amount (₹)"]

        validation_issues = []

        # ── Sheet 1: Summary ─────────────────────────────────────────────────
        ws_summary = wb.create_sheet("Summary")
        ws_summary.column_dimensions["A"].width = 25
        ws_summary.column_dimensions["B"].width = 40

        ws_summary.append(["BOQ Report Summary"])
        ws_summary["A1"].font = Font(name="Calibri", bold=True, size=16, color="003366")
        ws_summary.merge_cells("A1:B1")
        ws_summary.append([])
        ws_summary.append(["Project Name", getattr(parse_result, "project_name", "")])
        ws_summary.append(["Report Type", getattr(parse_result, "report_type", "")])
        ws_summary.append(["Source File", getattr(parse_result, "source_file", "")])
        ws_summary.append(["Total Pages", getattr(parse_result, "total_pages", "")])
        ws_summary.append(["PDF Type", getattr(parse_result, "pdf_type", "")])
        ws_summary.append(["OCR Mode", getattr(parse_result, "ocr_mode", "")])
        ws_summary.append([])
        vs = getattr(parse_result, "validation_summary", None)
        if vs:
            ws_summary.append(["Total Rows", vs.total_rows])
            ws_summary.append(["Valid Rows", vs.valid_rows])
            ws_summary.append(["Issues Found", vs.invalid_rows])
            ws_summary.append(["Grand Total (₹)", vs.grand_total_computed])

        # ── Sheet 2+: one sheet per top-level section ─────────────────────────
        for sec in getattr(parse_result, "sections", []):
            sheet_title = f"{sec.number} {sec.title}"[:31]  # Excel limit
            ws = wb.create_sheet(sheet_title)

            # Set column widths
            for i, w in enumerate(col_widths):
                ws.column_dimensions[chr(65 + i)].width = w

            # Title row
            ws.append([f"{sec.number} {sec.title}"])
            ws.row_dimensions[1].height = 22
            title_cell = ws.cell(1, 1)
            title_cell.font = title_font
            title_cell.fill = title_fill
            title_cell.alignment = center_align
            ws.merge_cells(f"A1:{chr(64 + len(headers))}1")

            # Header row
            ws.append(headers)
            ws.row_dimensions[2].height = 18
            for c_idx, h in enumerate(headers, 1):
                cell = ws.cell(2, c_idx)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = center_align
                cell.border = border

            row_num = 3
            row_num = self._write_section(
                ws, sec, row_num, section_font, section_fill,
                normal_font, subtotal_fill, error_fill, error_font,
                alt_fill, border, left_align, right_align, center_align,
                col_widths, validation_issues
            )

        # ── Sheet: Validation Issues ─────────────────────────────────────────
        if validation_issues:
            ws_issues = wb.create_sheet("Validation Issues")
            ws_issues.column_dimensions["A"].width = 10
            ws_issues.column_dimensions["B"].width = 40
            ws_issues.column_dimensions["C"].width = 20
            ws_issues.column_dimensions["D"].width = 14
            ws_issues.column_dimensions["E"].width = 14
            ws_issues.column_dimensions["F"].width = 10

            ws_issues.append(["Item No.", "Description", "Issue Type", "Expected (₹)", "Actual (₹)", "Error %"])
            for r in validation_issues:
                ws_issues.append(r)

        out = Path(output_path)
        out.parent.mkdir(parents=True, exist_ok=True)
        wb.save(str(out))
        return str(out.resolve())

    def _write_section(self, ws, section, row_num, section_font, section_fill,
                       normal_font, subtotal_fill, error_fill, error_font,
                       alt_fill, border, left_align, right_align, center_align,
                       col_widths, validation_issues, depth=0):
        """Recursively write a section and its children."""
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        import openpyxl

        n_cols = len(col_widths)

        # Section heading row
        ws.append([
            section.number,
            "  " * depth + section.title,
            "", "", "", ""
        ])
        ws.row_dimensions[row_num].height = 16
        for c_idx in range(1, n_cols + 1):
            cell = ws.cell(row_num, c_idx)
            cell.font = section_font
            cell.fill = section_fill
        ws.cell(row_num, 1).alignment = center_align
        ws.cell(row_num, 2).alignment = left_align
        row_num += 1

        # Item rows
        for i, item in enumerate(section.items):
            if item.is_header_row:
                continue

            if item.is_subtotal_row:
                ws.append([
                    "", item.description or "Subtotal",
                    "", "", "",
                    item.amount if item.amount is not None else ""
                ])
                ws.row_dimensions[row_num].height = 14
                for c_idx in range(1, n_cols + 1):
                    ws.cell(row_num, c_idx).fill = subtotal_fill
                row_num += 1
                continue

            qty = _safe_float(item.quantity)
            rate = _safe_float(item.rate)
            amt = _safe_float(item.amount)

            # Check for mismatch
            is_mismatch = False
            if qty is not None and rate is not None and amt is not None:
                expected = qty * rate
                if abs(expected - amt) / max(abs(amt), 1.0) > 0.02:
                    is_mismatch = True
                    validation_issues.append([
                        item.item_no, item.description, "amount_mismatch",
                        round(expected, 2), amt,
                        round(abs(expected - amt) / max(abs(amt), 1.0) * 100, 1)
                    ])

            ws.append([
                item.item_no,
                "  " * (depth + 1) + item.description,
                qty, item.unit, rate, amt
            ])
            ws.row_dimensions[row_num].height = 14

            row_fill = error_fill if is_mismatch else (None if i % 2 == 0 else None)
            for c_idx in range(1, n_cols + 1):
                cell = ws.cell(row_num, c_idx)
                if is_mismatch:
                    cell.fill = error_fill
                    cell.font = error_font
                else:
                    cell.font = normal_font
                if c_idx == 2:
                    cell.alignment = left_align
                elif c_idx in (3, 5, 6):
                    cell.alignment = right_align
                    if c_idx in (5, 6) and cell.value is not None:
                        cell.number_format = '#,##0.00'
                else:
                    cell.alignment = center_align

            row_num += 1

        # Recurse into children
        for child in section.children:
            row_num = self._write_section(
                ws, child, row_num, section_font, section_fill,
                normal_font, subtotal_fill, error_fill, error_font,
                None, border, left_align, right_align, center_align,
                col_widths, validation_issues, depth=depth + 1
            )

        # Section total row
        if section.subtotal is not None:
            ws.append(["", f"Total - {section.title}", "", "", "", section.subtotal])
            ws.row_dimensions[row_num].height = 15
            from openpyxl.styles import Font as _Font, PatternFill as _Fill
            for c_idx in range(1, n_cols + 1):
                cell = ws.cell(row_num, c_idx)
                cell.font = _Font(name="Calibri", bold=True, size=10)
                cell.fill = _Fill("solid", fgColor="FFF2CC")
            ws.cell(row_num, 6).number_format = '#,##0.00'
            row_num += 1

        return row_num
