"""Tests for Export modules."""

from __future__ import annotations

import json
import os
import sys
import tempfile
import pytest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", ".."))

from boqparser.table.table_reconstructor import BOQRow
from boqparser.hierarchy.hierarchy_detector import Section
from boqparser.validation.validator import ValidationSummary
from boqparser.parser import ParseResult


def make_test_parse_result() -> ParseResult:
    """Create a minimal ParseResult for export testing."""
    sec1 = Section(number="1", title="FOUNDATION WORKS", depth=0)
    sec1.items = [
        BOQRow(item_no="1.1.1", description="Earth Work for Footing", quantity=6035.75, unit="Cft", rate=11.0, amount=66393.25, page_num=1),
        BOQRow(item_no="1.1.2", description="PCC 1:4:8", quantity=150.5, unit="CuM", rate=4500.0, amount=677250.0, page_num=1),
    ]
    sec1.subtotal = 743643.25

    sec1_1 = Section(number="1.1", title="Earth Work", depth=1)
    sec1_1.items = [
        BOQRow(item_no="1.1.3", description="Backfilling", quantity=200, unit="Cft", rate=5.0, amount=1000.0, page_num=2),
    ]
    sec1.children = [sec1_1]

    all_rows = sec1.items + sec1_1.items

    vs = ValidationSummary(
        total_rows=3,
        valid_rows=3,
        invalid_rows=0,
        grand_total_computed=744643.25,
    )

    return ParseResult(
        source_file="/test/boq_report.pdf",
        project_name="PROPOSED RESIDENTIAL BUILDING",
        report_type="BOQ",
        pdf_type="digital",
        ocr_mode="fallback",
        total_pages=5,
        sections=[sec1],
        all_rows=all_rows,
        validation_summary=vs,
        metadata={"test": True},
    )


class TestJSONExporter:
    def test_export_to_file(self):
        from boqparser.export import JSONExporter
        result = make_test_parse_result()
        with tempfile.NamedTemporaryFile(suffix=".json", delete=False) as tf:
            path = tf.name
        try:
            JSONExporter().export(result, path)
            with open(path, encoding="utf-8") as f:
                data = json.load(f)
            assert data["project_name"] == "PROPOSED RESIDENTIAL BUILDING"
            assert data["report_type"] == "BOQ"
            assert "sections" in data
            assert "all_items" in data
            assert "validation" in data
        finally:
            os.unlink(path)

    def test_to_string(self):
        from boqparser.export import JSONExporter
        result = make_test_parse_result()
        json_str = JSONExporter().to_string(result)
        data = json.loads(json_str)
        assert data["project_name"] == "PROPOSED RESIDENTIAL BUILDING"
        assert len(data["all_items"]) == 3


class TestCSVExporter:
    def test_export_to_file(self):
        from boqparser.export import CSVExporter
        result = make_test_parse_result()
        with tempfile.NamedTemporaryFile(suffix=".csv", delete=False, mode="w") as tf:
            path = tf.name
        try:
            CSVExporter().export(result, path)
            with open(path, encoding="utf-8-sig") as f:
                content = f.read()
            assert "section_path" in content
            assert "item_no" in content
            assert "Earth Work for Footing" in content
        finally:
            os.unlink(path)

    def test_to_string(self):
        from boqparser.export import CSVExporter
        result = make_test_parse_result()
        csv_str = CSVExporter().to_string(result)
        assert "section_path" in csv_str
        assert "1.1.1" in csv_str


class TestMarkdownExporter:
    def test_export_to_file(self):
        from boqparser.export import MarkdownExporter
        result = make_test_parse_result()
        with tempfile.NamedTemporaryFile(suffix=".md", delete=False, mode="w") as tf:
            path = tf.name
        try:
            MarkdownExporter().export(result, path)
            with open(path, encoding="utf-8") as f:
                content = f.read()
            assert "PROPOSED RESIDENTIAL BUILDING" in content
            assert "FOUNDATION WORKS" in content
            assert "Earth Work for Footing" in content
            assert "|" in content  # Markdown table
        finally:
            os.unlink(path)

    def test_to_string(self):
        from boqparser.export import MarkdownExporter
        result = make_test_parse_result()
        md = MarkdownExporter().to_string(result)
        assert "# PROPOSED RESIDENTIAL BUILDING" in md
        assert "## 1 FOUNDATION WORKS" in md


class TestDBExporter:
    def test_export_to_sqlite(self):
        from boqparser.export import DBExporter
        result = make_test_parse_result()
        with tempfile.NamedTemporaryFile(suffix=".db", delete=False) as tf:
            path = tf.name
        try:
            DBExporter().export(result, path)
            import sqlite3
            conn = sqlite3.connect(path)
            cur = conn.cursor()
            cur.execute("SELECT name FROM projects")
            rows = cur.fetchall()
            assert len(rows) == 1
            assert rows[0][0] == "PROPOSED RESIDENTIAL BUILDING"

            cur.execute("SELECT COUNT(*) FROM items")
            count = cur.fetchone()[0]
            assert count >= 2

            cur.execute("SELECT COUNT(*) FROM sections")
            sec_count = cur.fetchone()[0]
            assert sec_count >= 1
            conn.close()
        finally:
            os.unlink(path)


class TestExcelExporter:
    def test_export_to_excel(self):
        try:
            import openpyxl
        except ImportError:
            pytest.skip("openpyxl not installed")

        from boqparser.export import ExcelExporter
        result = make_test_parse_result()
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tf:
            path = tf.name
        try:
            ExcelExporter().export(result, path)
            import openpyxl
            wb = openpyxl.load_workbook(path)
            # Should have at least Summary + one section sheet
            assert "Summary" in wb.sheetnames
            assert len(wb.sheetnames) >= 2
        finally:
            os.unlink(path)
